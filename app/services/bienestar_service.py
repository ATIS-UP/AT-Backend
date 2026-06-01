"""Service for Bienestar TCBU data: query, CSV bulk-load, template generation."""
import csv
import io
from typing import Optional

from fastapi import UploadFile
from sqlalchemy.orm import Session

from app.exceptions import ValidationError
from app.models.bienestar import BienestarRegistro, SERVICIOS_BIENESTAR
from app.utils.audit import AuditService


def _periodo_sort_key(p: str) -> tuple[int, int]:
    """Sort key for period strings like '2021-1'. Works for any future year."""
    try:
        year, sem = p.split("-")
        return (int(year), int(sem))
    except (ValueError, AttributeError):
        return (9999, 9)


# CSV column header → canonical service name
_CSV_ALIAS: dict[str, str] = {
    "cultural": "Cultural",
    "inclusión": "Inclusión",
    "inclusion": "Inclusión",
    "act. física": "Act. Física",
    "act.física": "Act. Física",
    "act.fisica": "Act. Física",
    "actividad física": "Act. Física",
    "salud": "Salud",
    "socioeconómica": "Socioeconómica",
    "socioeconomica": "Socioeconómica",
    "socioecon": "Socioeconómica",
    "socioec": "Socioeconómica",
    "espiritual": "Espiritual",
    "psicológica": "Psicológica",
    "psicologica": "Psicológica",
    "psicol": "Psicológica",
    "odontología": "Odontología",
    "odontologia": "Odontología",
    "odont": "Odontología",
    "alimentación": "Alimentación",
    "alimentacion": "Alimentación",
    "aliment": "Alimentación",
    "simup": "SIMUP",
}


class BienestarService:
    def __init__(self, db: Session):
        self.db = db

    # ── TCBU query ────────────────────────────────────────────────────────────

    def get_tcbu(
        self,
        periodo_inicio: Optional[str] = None,
        periodo_fin: Optional[str] = None,
    ) -> dict:
        """Return TCBU data pivoted for a multi-line chart.

        Response shape:
        {
          "periodos": ["2021-1", ...],
          "servicios": ["Cultural", ...],
          "series": [
            {"periodo": "2021-1", "Cultural": 77, "Inclusión": 119, ...},
            ...
          ]
        }
        """
        rows = self.db.query(BienestarRegistro).all()
        present_periods = sorted({r.periodo for r in rows}, key=_periodo_sort_key)

        if periodo_inicio:
            present_periods = [p for p in present_periods if p >= periodo_inicio]
        if periodo_fin:
            present_periods = [p for p in present_periods if p <= periodo_fin]

        # build pivot: {periodo: {servicio: cantidad}}
        pivot: dict[str, dict[str, int]] = {p: {} for p in present_periods}
        for r in rows:
            if r.periodo in pivot:
                pivot[r.periodo][r.servicio] = r.cantidad

        series = [
            {"periodo": p, **{s: pivot[p].get(s, 0) for s in SERVICIOS_BIENESTAR}}
            for p in present_periods
        ]

        return {
            "periodos": present_periods,
            "servicios": SERVICIOS_BIENESTAR,
            "series": series,
        }

    # ── CSV bulk load ─────────────────────────────────────────────────────────

    def carga_csv(self, file: UploadFile, usuario_id: str, ip: str) -> dict:
        """Parse wide-format CSV/XLSX and upsert records.

        Wide format (one row per period):
          periodo, Cultural, Inclusión, Act. Física, Salud, ...
          2025-1,  51,       523,       146,          228, ...
        """
        ext = self._ext(file.filename)
        if ext not in {".csv", ".xlsx"}:
            raise ValidationError("Solo se aceptan archivos .csv o .xlsx")

        content = file.file.read()
        rows = self._parse_csv(content) if ext == ".csv" else self._parse_xlsx(content)

        if not rows:
            raise ValidationError("El archivo está vacío o no tiene datos válidos")

        headers = list(rows[0].keys())
        if "periodo" not in headers:
            raise ValidationError('Columna "periodo" requerida', fields={"archivo": "falta columna periodo"})

        # map headers to canonical service names
        header_map: dict[str, str] = {}
        for h in headers:
            if h == "periodo":
                continue
            canonical = _CSV_ALIAS.get(h.lower())
            if canonical:
                header_map[h] = canonical

        if not header_map:
            raise ValidationError(
                "No se encontró ninguna columna de servicio reconocida",
                fields={"archivo": f"columnas detectadas: {', '.join(headers)}"},
            )

        # collect valid periods from the file first
        periodos_en_archivo = {
            str(r.get("periodo", "")).strip()
            for r in rows
            if str(r.get("periodo", "")).strip()
        }

        # single query to load all existing rows for those periods
        existing_records: dict[tuple[str, str], BienestarRegistro] = {
            (rec.periodo, rec.servicio): rec
            for rec in self.db.query(BienestarRegistro).filter(
                BienestarRegistro.periodo.in_(periodos_en_archivo)
            ).all()
        }

        insertados = 0
        actualizados = 0
        errores: list[dict] = []

        for idx, row in enumerate(rows, start=2):
            periodo = str(row.get("periodo", "")).strip()
            if not periodo:
                errores.append({"fila": idx, "error": "periodo vacío"})
                continue

            for raw_col, servicio in header_map.items():
                val_str = str(row.get(raw_col, "0")).strip()
                try:
                    cantidad = int(float(val_str)) if val_str else 0
                except (ValueError, TypeError):
                    errores.append({"fila": idx, "error": f"valor inválido en {raw_col}: {val_str!r}"})
                    continue

                existing = existing_records.get((periodo, servicio))
                if existing:
                    existing.cantidad = cantidad
                    existing.uploaded_by = usuario_id
                    actualizados += 1
                else:
                    new_rec = BienestarRegistro(
                        periodo=periodo,
                        servicio=servicio,
                        cantidad=cantidad,
                        uploaded_by=usuario_id,
                    )
                    self.db.add(new_rec)
                    existing_records[(periodo, servicio)] = new_rec
                    insertados += 1

        self.db.commit()

        AuditService.log(
            db=self.db,
            usuario_id=usuario_id,
            accion="CARGA_BIENESTAR",
            entidad="BienestarRegistro",
            detalles={"insertados": insertados, "actualizados": actualizados, "errores": len(errores)},
            ip=ip,
        )

        return {
            "insertados": insertados,
            "actualizados": actualizados,
            "errores": errores,
            "total_filas": len(rows),
        }

    # ── CSV template ──────────────────────────────────────────────────────────

    def generar_plantilla_csv(self) -> bytes:
        """Return a downloadable CSV template with headers and one example row."""
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["periodo"] + SERVICIOS_BIENESTAR)
        writer.writerow(["2026-1"] + [0] * len(SERVICIOS_BIENESTAR))
        return output.getvalue().encode("utf-8-sig")

    # ── helpers ───────────────────────────────────────────────────────────────

    @staticmethod
    def _ext(filename: Optional[str]) -> str:
        if not filename:
            return ""
        idx = filename.rfind(".")
        return filename[idx:].lower() if idx != -1 else ""

    @staticmethod
    def _parse_csv(content: bytes) -> list[dict]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        reader = csv.DictReader(io.StringIO(text))
        return [
            {k.strip().lower(): v.strip() if v else "" for k, v in row.items() if k}
            for row in reader
        ]

    @staticmethod
    def _parse_xlsx(content: bytes) -> list[dict]:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValidationError("Instale openpyxl para soporte XLSX")

        wb = load_workbook(io.BytesIO(content), read_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = [str(h).strip().lower() if h else "" for h in next(rows_iter)]
        except StopIteration:
            return []

        result = []
        for row_vals in rows_iter:
            d = {
                headers[i]: (str(v).strip() if v is not None else "")
                for i, v in enumerate(row_vals)
                if i < len(headers) and headers[i]
            }
            result.append(d)
        return result
